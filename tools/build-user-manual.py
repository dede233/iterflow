"""Export the canonical user manual using the artifact runtime's python-docx.

Usage: <bundled-python> tools/build-user-manual.py --output <manual.docx>
Render the output and inspect every page before delivering it.
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def plain(text):
    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text).replace("`", "")


def build(source, output):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    for name, size, bold in (
        ("Normal", 11, False),
        ("Title", 25, True),
        ("Heading 1", 17, True),
        ("List Number", 11, False),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for attr in list(fonts.attrib):
            if "theme" in attr.lower():
                del fonts.attrib[attr]
        fonts.set(qn("w:eastAsia"), "Heiti SC")
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.2
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    doc.core_properties.title = "迭程 IterFlow 操作手册"
    doc.core_properties.author = "IterFlow"
    doc.core_properties.subject = "用户操作与管理员日常维护"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("迭程 IterFlow  |  ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = source.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.fullmatch(r"[| :\-]+", lines[i]):
                    rows.append(
                        [plain(c.strip()) for c in lines[i].strip("|").split("|")]
                    )
                i += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.autofit = False
            widths = (2.05, 4.85)
            for column, width in zip(table.columns, widths):
                column.width = Inches(width)
            for row_index, values in enumerate(rows):
                row = table.add_row()
                properties = row._tr.get_or_add_trPr()
                properties.append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    properties.append(OxmlElement("w:tblHeader"))
                for col, (cell, value) in enumerate(zip(row.cells, values)):
                    cell.width = Inches(widths[col])
                    cell.text = value
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    props = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for side in ("top", "left", "bottom", "right"):
                        edge = OxmlElement("w:" + side)
                        for key, val in (
                            ("val", "single"),
                            ("sz", "4"),
                            ("color", "D9D9D9"),
                        ):
                            edge.set(qn("w:" + key), val)
                        borders.append(edge)
                    props.append(borders)
                    margins = OxmlElement("w:tcMar")
                    for side in ("top", "left", "bottom", "right"):
                        margin = OxmlElement("w:" + side)
                        margin.set(qn("w:w"), "100")
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    props.append(margins)
                    if row_index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "E8EEF5")
                        props.append(shading)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(2)
                        paragraph.paragraph_format.line_spacing = 1.15
                        for run in paragraph.runs:
                            run.font.size = Pt(10)
                            run.font.bold = row_index == 0
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if line.startswith("# "):
            doc.add_paragraph(plain(line[2:]), "Title")
        elif line.startswith("## "):
            doc.add_heading(plain(line[3:]), 1)
        elif line.strip():
            # Keep the explicit fifteen-step numbering identical to the source.
            paragraph = doc.add_paragraph(plain(line))
            paragraph.paragraph_format.widow_control = True
        i += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(Path(__file__).resolve().parents[1] / "docs/user-manual.md", args.output)
    print(args.output)
